from __future__ import annotations
import os
import datetime as dt
from collections import defaultdict
import openpyxl
import pandas as pd
from engine import MessageEngine
from coordinator_parser import clean_phone_number, value_to_text
from coordinator_logic import _arabic_day
from db_manager import get_all_centers
from file_errors import FileReadError, EmptyFileError, MissingColumnsError

OUTPUT_MESSAGES_FILE = "center_messages_output.xlsx"
OUTPUT_LOG_FILE = "center_whatsapp_log.xlsx"

DEFAULT_CENTER_TEMPLATE = (
    "برجاء تأكيد الحجز للمواعيد التالية في {center_name}:\n\n"
    "{sections}"
)

def load_sessions_for_centers(file_path: str) -> list[dict]:
    """Parse the sessions file specifically for Centers requirements."""
    filename = os.path.basename(file_path)
    ext = os.path.splitext(file_path)[1].lower()

    if ext not in {".xlsx", ".xlsm"}:
        raise FileReadError(filename, "Only .xlsx and .xlsm files are supported.")

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as exc:
        raise FileReadError(filename, str(exc)) from exc

    # Find a sheet with center, day, time, group, room
    target_ws = None
    headers = []
    canonical_map = {}
    
    needed_cols = {
        "center_name": {"centername", "center", "venue", "venuename", "location"},
        "day": {"day", "sessionday"},
        "time": {"time", "sessiontime", "starttime"},
        "group_code": {"groupcode", "group", "groupid"},
        "room_code": {"roomcode", "room", "roomid"},
    }

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue
        h = [value_to_text(v) for v in rows[0]]
        
        c_map = {}
        for idx, header in enumerate(h):
            norm = value_to_text(header).lower().replace(" ", "").replace("_", "")
            for canonical, variants in needed_cols.items():
                if norm in variants or norm == canonical:
                    c_map[canonical] = idx
        
        if set(needed_cols.keys()).issubset(c_map.keys()):
            target_ws = ws
            headers = h
            canonical_map = c_map
            break

    if not target_ws:
        raise MissingColumnsError(
            missing=list(needed_cols.keys()),
            detected=[],
            accepted_hints=needed_cols
        )

    rows = list(target_ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise EmptyFileError(filename)

    sessions = []
    for row in rows[1:]:
        session = {}
        has_val = False
        for k, idx in canonical_map.items():
            val = value_to_text(row[idx] if idx < len(row) else "")
            session[k] = val
            if val: has_val = True
        if has_val and session["center_name"]:
            sessions.append(session)

    return sessions

def generate_center_messages(sessions: list[dict], template_text: str = None):
    """
    Groups sessions by center, matches with DB to get POC info,
    calculates slots and max rooms, and generates messages.
    """
    db_centers = get_all_centers()
    # Create lookup by center name (case-insensitive, stripped)
    center_lookup = {c["center_name"].strip().casefold(): c for c in db_centers}

    # Group by center
    by_center = defaultdict(list)
    for s in sessions:
        by_center[s["center_name"].strip()].append(s)

    messages = []
    skipped_logs = []

    for center_name, center_sessions in by_center.items():
        db_info = center_lookup.get(center_name.casefold())
        if not db_info:
            skipped_logs.append({
                "Timestamp": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "Name": center_name,
                "Phone": "",
                "Status": "skipped",
                "Error": "Center not found in Database. Please upload it in Centers List first."
            })
            continue
        
        poc_name = db_info.get("poc", "").strip() or "مسؤول السنتر"
        phone = clean_phone_number(db_info.get("telephone", ""))
        
        if not phone:
            skipped_logs.append({
                "Timestamp": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "Name": center_name,
                "Phone": "",
                "Status": "skipped",
                "Error": "Center has no valid telephone number in DB."
            })
            continue

        # Group by day
        by_day = defaultdict(list)
        for s in center_sessions:
            by_day[s["day"]].append(s)

        sections = []
        total_slots = 0
        max_rooms_overall = 0

        for day, day_sessions in sorted(by_day.items()):
            display_day = _arabic_day(day)
            
            # Group by time to find slots and rooms needed
            by_time = defaultdict(list)
            for s in day_sessions:
                by_time[s["time"]].append(s)
            
            day_slots = len(by_time)
            total_slots += day_slots
            
            # Max rooms needed this day = max concurrent sessions in any time slot
            day_max_rooms = 0
            slot_details = []
            for time_slot, time_sessions in sorted(by_time.items()):
                unique_groups = set(s["group_code"] for s in time_sessions if s["group_code"])
                rooms_needed = len(unique_groups)
                if rooms_needed > day_max_rooms:
                    day_max_rooms = rooms_needed
                slot_details.append(f"  • فترة {time_slot}: (عدد {rooms_needed} سيشن/قاعات)")
            
            if day_max_rooms > max_rooms_overall:
                max_rooms_overall = day_max_rooms

            slots_joined = "\n".join(slot_details)
            sections.append(
                f"يوم {display_day}:\n"
                f"- تفاصيل الفترات:\n{slots_joined}\n"
                f"- إجمالي الفترات: {day_slots}"
            )

        template = template_text or DEFAULT_CENTER_TEMPLATE
        message = template.replace("{poc_name}", poc_name)
        message = message.replace("{center_name}", center_name)
        message = message.replace("{sections}", "\n\n".join(sections))

        messages.append({
            "Center": center_name,
            "POC": poc_name,
            "Phone": phone,
            "Total Slots": total_slots,
            "Max Rooms Needed": max_rooms_overall,
            "Message": message.strip()
        })

    return messages, skipped_logs

def save_outputs(messages, skipped_logs, output_dir=None):
    output_dir = output_dir or os.getcwd()
    os.makedirs(output_dir, exist_ok=True)
    
    msg_path = os.path.join(output_dir, OUTPUT_MESSAGES_FILE)
    log_path = os.path.join(output_dir, OUTPUT_LOG_FILE)
    
    if messages:
        df_msg = pd.DataFrame(messages)
        df_msg.to_excel(msg_path, index=False)
        
    if skipped_logs:
        df_log = pd.DataFrame(skipped_logs)
        df_log.to_excel(log_path, index=False)
        
    return {"messages": msg_path, "log": log_path}
