from __future__ import annotations
import csv
import datetime as dt
import os
import re
import openpyxl
from file_errors import (
    FileReadError, EmptyFileError, MissingColumnsError, CoordinatorDirectoryError,
)

CANONICAL_COLUMNS = {
    "group_code": {"groupcode", "group", "groupid"},
    "center_code": {"centercode", "centercodeid"},
    "room_code": {"roomcode", "room", "roomid"},
    "area": {"area", "region", "zone"},
    "city": {"city", "governorate", "governate"},
    "center_name": {"centername", "center", "venue", "venuename", "location"},
    "day": {"day", "sessionday"},
    "time": {"time", "sessiontime", "starttime"},
    "main_coordinator": {"maincoordinator", "coordinator", "main", "mainname"},
    "main_phone": {"phonenumber", "phone", "mainphone", "mainphonenumber", "coordinatorphone"},
    "main_code": {"code", "maincode", "coordinatorcode"},
    "backup_coordinator": {"backupcoordinator", "backup", "backupname"},
    "backup_phone": {"phonenumber2", "phone2", "backupphone", "backupphonenumber"},
    "backup_code": {"code2", "backupcode"},
    "no_students": {"nostudents", "students", "studentcount", "noofstudents"},
}
BASE_REQUIRED_COLUMNS = {"group_code", "room_code", "day", "time", "center_name"}
ROLE_NAME_COLUMNS = ("main_coordinator", "backup_coordinator")
ROLE_PHONE_COLUMNS = ("main_phone", "backup_phone")
COORDINATOR_DIRECTORY_COLUMNS = {
    "coordinator_name": {"coordinatorname", "coordinator", "name", "fullname", "coordinatorfullname", "اسمالمنسق", "الاسم", "اسم"},
    "phone": {"phonenumber", "phone", "mobile", "mobilenumber", "whatsapp", "whatsappnumber", "رقمالهاتف", "رقمالموبايل", "الموبايل", "الهاتف"},
    "gender": {"gender", "sex", "type", "malefemale", "النوع", "الجنس"},
}

def load_data(file_path: str, coordinator_file_path: str | None = None) -> dict:
    """Read the sessions file and optionally merge a separate coordinator phone directory."""
    filename = os.path.basename(file_path)
    ext = os.path.splitext(file_path)[1].lower()

    if ext not in {".xlsx", ".xlsm", ".csv"}:
        raise FileReadError(
            filename,
            f"Unsupported file type '{ext}'. Only .xlsx, .xlsm, and .csv files are accepted.",
        )

    if not os.path.isfile(file_path):
        raise FileReadError(filename, f"File not found at:\n{file_path}")

    if ext in {".xlsx", ".xlsm"}:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as exc:
            raise FileReadError(filename, str(exc)) from exc

        sheets = []
        coordinator_directory = {}
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            sheet_data = {
                "name": ws.title,
                "headers": [value_to_text(v) for v in rows[0]],
                "rows": [list(r) for r in rows[1:]],
            }
            sheets.append(sheet_data)
            coordinator_directory.update(_read_coordinator_directory(sheet_data["headers"], sheet_data["rows"]))

        if not sheets:
            raise EmptyFileError(filename)

        session_sheet = _choose_session_sheet(sheets, wb.active.title)
        if coordinator_file_path:
            coordinator_directory.update(load_coordinator_directory(coordinator_file_path))
        return {
            "headers": session_sheet["headers"],
            "rows": session_sheet["rows"],
            "sheet_name": session_sheet["name"],
            "coordinator_directory": coordinator_directory,
        }

    # CSV path
    rows = []
    for encoding in ("utf-8-sig", "utf-8", "cp1256"):
        try:
            with open(file_path, newline="", encoding=encoding) as f:
                reader = csv.reader(f)
                rows = list(reader)
            break
        except UnicodeDecodeError:
            rows = []
    if not rows:
        raise EmptyFileError(filename)
    coordinator_directory = load_coordinator_directory(coordinator_file_path) if coordinator_file_path else {}
    return {"headers": [value_to_text(v) for v in rows[0]], "rows": rows[1:], "coordinator_directory": coordinator_directory}

def load_coordinator_directory(file_path: str) -> dict:
    """Read a coordinator data file with Coordinator name + phone number columns."""
    filename = os.path.basename(file_path)
    ext = os.path.splitext(file_path)[1].lower()

    if ext not in {".xlsx", ".xlsm", ".csv"}:
        raise FileReadError(
            filename,
            f"Unsupported file type '{ext}'. Only .xlsx, .xlsm, and .csv files are accepted.",
        )

    if ext in {".xlsx", ".xlsm"}:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as exc:
            raise FileReadError(filename, str(exc)) from exc

        directory = {}
        all_headers = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [value_to_text(v) for v in rows[0]]
            all_headers.extend(h for h in headers if h)
            directory.update(_read_coordinator_directory(headers, [list(r) for r in rows[1:]]))
        if not directory:
            raise CoordinatorDirectoryError(filename, detected_headers=all_headers or None)
        return directory

    # CSV path
    rows = []
    for encoding in ("utf-8-sig", "utf-8", "cp1256"):
        try:
            with open(file_path, newline="", encoding=encoding) as f:
                rows = list(csv.reader(f))
            break
        except UnicodeDecodeError:
            rows = []
    if not rows:
        raise EmptyFileError(filename)
    headers = [value_to_text(v) for v in rows[0]]
    directory = _read_coordinator_directory(headers, rows[1:])
    if not directory:
        detected = [h for h in headers if h]
        raise CoordinatorDirectoryError(filename, detected_headers=detected or None)
    return directory

def validate_columns(data: dict) -> list[dict]:
    """Map supported column-name variations into the canonical coordinator schema."""
    headers = data["headers"]
    coordinator_directory = data.get("coordinator_directory", {})
    column_map = _build_column_map(headers)
    missing = sorted(BASE_REQUIRED_COLUMNS - set(column_map))
    has_role_names = any(col in column_map for col in ROLE_NAME_COLUMNS)
    has_phone_source = any(col in column_map for col in ROLE_PHONE_COLUMNS) or bool(coordinator_directory)

    if missing or not has_role_names or not has_phone_source:
        all_missing = list(missing)
        extra_notes = []
        if not has_role_names:
            all_missing.extend(col for col in ROLE_NAME_COLUMNS if col not in column_map)
            extra_notes.append(
                "No coordinator name columns were found. Add a 'Main Coordinator' "
                "and/or 'Backup Coordinator' column to the sessions sheet."
            )
        if not has_phone_source:
            all_missing.extend(col for col in ROLE_PHONE_COLUMNS if col not in column_map)
            extra_notes.append(
                "No phone columns were found. Either add phone columns to the sessions sheet, "
                "or provide a separate coordinators data file with 'Coordinator name' and "
                "'phone number' columns."
            )

        detected = [h for h in headers if h]
        # Build accepted-name hints only for actually-missing columns
        hints = {}
        for col in all_missing:
            if col in CANONICAL_COLUMNS:
                hints[col] = CANONICAL_COLUMNS[col]

        raise MissingColumnsError(
            missing=all_missing,
            detected=detected if detected else None,
            accepted_hints=hints if hints else None,
        )

    canonical_rows = []
    for source_row in data["rows"]:
        row = {key: "" for key in CANONICAL_COLUMNS}
        for key, idx in column_map.items():
            row[key] = value_to_text(source_row[idx] if idx < len(source_row) else "")
        if not any(row.values()):
            continue
        main_directory_entry = coordinator_directory.get(_person_key(row["main_coordinator"]), "")
        backup_directory_entry = coordinator_directory.get(_person_key(row["backup_coordinator"]), "")
        # Use clean_phone_number to check validity before falling back to directory.
        # This prevents junk values like "0" from blocking the directory lookup.
        if not clean_phone_number(row["main_phone"]):
            row["main_phone"] = _directory_phone(main_directory_entry) or row["main_phone"]
        if not clean_phone_number(row["backup_phone"]):
            row["backup_phone"] = _directory_phone(backup_directory_entry) or row["backup_phone"]
        row["main_gender"] = _directory_gender(main_directory_entry)
        row["backup_gender"] = _directory_gender(backup_directory_entry)
        row["area"] = row["area"] or "غير محدد"
        row["city"] = row["city"] or "غير محدد"
        canonical_rows.append(row)

    if not canonical_rows:
        sheet_name = data.get("sheet_name", "unknown")
        raise EmptyFileError(
            f"sheet '{sheet_name}'"
        )
    return canonical_rows

def clean_phone_number(phone) -> str:
    """Clean and validate Egyptian WhatsApp phone numbers."""
    text = value_to_text(phone)
    if not text:
        return ""
    text = text.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789"))
    text = re.sub(r"\.0+$", "", text)
    text = re.sub(r"[\s\-\(\)]", "", text)
    text = "".join(ch for ch in text if ch.isdigit() or ch == "+")

    if text.startswith("00"):
        text = "+" + text[2:]
    elif text.startswith("0"):
        text = "+20" + text[1:]
    elif text.startswith("20"):
        text = "+20" + text[2:]
    elif re.fullmatch(r"1[0125]\d{8}", text):
        text = "+20" + text

    if re.fullmatch(r"\+201[0125]\d{8}", text):
        return text
    return ""

def value_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.date() == dt.date(1899, 12, 30):
            return value.strftime("%I:%M %p")
        return value.strftime("%Y-%m-%d %I:%M %p").strip()
    if isinstance(value, dt.time):
        return value.strftime("%I:%M %p")
    text = str(value).strip()
    return re.sub(r"\.0$", "", text)

def _build_column_map(headers):
    used = set()
    column_map = {}
    for idx, header in enumerate(headers):
        normalized = _normalize_header(header)
        if not normalized:
            continue
        for canonical, variants in CANONICAL_COLUMNS.items():
            if canonical in column_map:
                continue
            if normalized in variants or normalized == canonical:
                column_map[canonical] = idx
                used.add(idx)
                break

    # Helpful fallback for spreadsheets that duplicate "phone number" or "Code".
    phone_indexes = [i for i, h in enumerate(headers) if _normalize_header(h) in {"phonenumber", "phone"} and i not in used]
    if "backup_phone" not in column_map and phone_indexes:
        column_map["backup_phone"] = phone_indexes[0]
    code_indexes = [i for i, h in enumerate(headers) if _normalize_header(h) == "code" and i not in used]
    if "backup_code" not in column_map and code_indexes:
        column_map["backup_code"] = code_indexes[0]
    return column_map

def _choose_session_sheet(sheets, active_title):
    candidates = []
    for sheet in sheets:
        column_map = _build_column_map(sheet["headers"])
        has_required = BASE_REQUIRED_COLUMNS <= set(column_map)
        has_role_name = any(col in column_map for col in ROLE_NAME_COLUMNS)
        score = sum(1 for col in CANONICAL_COLUMNS if col in column_map)
        if has_required and has_role_name:
            candidates.append((sheet["name"] == active_title, score, sheet))

    if candidates:
        candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
        return candidates[0][2]

    active_sheet = next((sheet for sheet in sheets if sheet["name"] == active_title), None)
    return active_sheet or sheets[0]

def _read_coordinator_directory(headers, rows):
    column_map = _build_directory_column_map(headers)
    if {"coordinator_name", "phone"} - set(column_map):
        return {}

    directory = {}
    for row in rows:
        name = value_to_text(row[column_map["coordinator_name"]] if column_map["coordinator_name"] < len(row) else "")
        phone = value_to_text(row[column_map["phone"]] if column_map["phone"] < len(row) else "")
        gender = value_to_text(row[column_map["gender"]] if column_map.get("gender", 999999) < len(row) else "")
        if name and phone:
            directory[_person_key(name)] = {"phone": phone, "gender": gender}
    return directory

def _build_directory_column_map(headers):
    column_map = {}
    for idx, header in enumerate(headers):
        normalized = _normalize_header(header)
        for canonical, variants in COORDINATOR_DIRECTORY_COLUMNS.items():
            if canonical not in column_map and normalized in variants:
                column_map[canonical] = idx
                break
    return column_map

def _normalize_header(header) -> str:
    text = value_to_text(header).lower()
    text = re.sub(r"[\u064b-\u065f\u0670ـ]", "", text)
    text = text.translate(str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه"}))
    return re.sub(r"[^a-z0-9\u0600-\u06ff]+", "", text)

def _person_key(name) -> str:
    return re.sub(r"\s+", " ", value_to_text(name).strip()).casefold()

def _directory_phone(entry):
    if isinstance(entry, dict):
        return entry.get("phone", "")
    return entry or ""

def _directory_gender(entry):
    if isinstance(entry, dict):
        return entry.get("gender", "")
    return ""
