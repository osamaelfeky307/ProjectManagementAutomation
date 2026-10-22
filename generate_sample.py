"""
Generate a sample Excel file for testing the Project Management Automation app.
Run this once to create sample_sessions.xlsx.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sessions"

# Headers
headers = ["Instructor Name", "Phone Number", "Group Code", "Day", "Time"]
header_fill = PatternFill(start_color="1a73e8", end_color="1a73e8", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=12)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')
    cell.border = thin_border

# Sample data
data = [
    ["Ahmad Al-Rashid",    "+966501234567", "GRP-7A", "Friday",   "09:00 AM"],
    ["Fatima Hassan",      "+966502345678", "GRP-7A", "Friday",   "09:00 AM"],
    ["Khaled Mansour",     "+966503456789", "GRP-7B", "Saturday", "09:00 AM"],
    ["Sara Al-Otaibi",     "+966504567890", "GRP-8A", "Friday",   "12:00 PM"],
    ["Mohammed Al-Harbi",  "+966505678901", "GRP-8A", "Saturday", "12:00 PM"],
    ["Noura Abdulaziz",    "+966506789012", "GRP-8B", "Friday",   "03:00 PM"],
]

for row_idx, row_data in enumerate(data, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

# Add notes for accepted formats
notes = [
    "⚠️ IMPORTANT FORMATTING NOTES ⚠️",
    "- Day: Single days (e.g., 'Friday', 'Saturday') or combined days (e.g., 'Friday/Saturday').",
    "- Time: Standard 12-hour or 24-hour formats (e.g., '09:00 AM', '03:00 PM', '15:00').",
    "- Phone Number: MUST include country code (e.g., Egyptian: '+201001234567' or '00201001234567')."
]

note_start_row = len(data) + 4
for i, note in enumerate(notes):
    row_idx = note_start_row + i
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(headers))
    cell = ws.cell(row=row_idx, column=1, value=note)
    if i == 0:
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[row_idx].height = 25
    else:
        cell.font = Font(bold=True, color="1E3A8A", size=11)
        cell.fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
        cell.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[row_idx].height = 22

# Adjust column widths
for col in range(1, len(headers) + 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20

wb.save("sample_sessions.xlsx")
print("sample_sessions.xlsx created successfully!")
