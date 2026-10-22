"""
Group Invite Logic - builds and sends WhatsApp group-invite messages
to students loaded from an Excel sheet.

Anti-ban measures:
  Personalized message per student (name placeholder filled in)
  Random delay 45-90s between sends (configurable)
  Daily batch limit default 25
  Progress persistence: resume next day without re-sending
"""
from __future__ import annotations
import datetime as dt, json, os, random, time, threading, re
import openpyxl
from engine import MessageEngine
from template_utils import render_template, get_first_name

GROUP_INVITE_TEMPLATES_FILE = os.path.join(os.path.dirname(__file__), "group_invite_templates.json")
GROUP_INVITE_PROGRESS_FILE  = os.path.join(os.path.dirname(__file__), "group_invite_progress.json")
GROUP_INVITE_LOG_FILE       = "group_invite_log.xlsx"

DEFAULT_GROUP_INVITE_TEMPLATES = {
    "1. Coaching Welcome (Arabic)": (
        "\u0635\u0628\u0627\u062d \u0627\u0644\u062e\u064a\u0631 \u064a\u0627 {name}\u060c \u0628\u0646\u0631\u062d\u0628 \u0628\u064a\u0643\u0645 \u0637\u0644\u0627\u0628 {program} \u0641\u0649 \u0645\u0631\u062d\u0644\u0629 \u0627\u0644\u0640 *{stage}* \ud83c\udf89\n"
        "\u0645\u0639\u0627\u0643\u0645 {coordinator_name} \u0627\u0644\u0645\u0646\u0633\u0642 \u0627\u0644\u062e\u0627\u0635 \u0628\u0627\u0644\u062c\u0631\u0648\u0628\u060c \u0648\u0647\u0643\u0648\u0646 \u0645\u0639\u0627\u0643\u0645 \u062e\u0644\u0627\u0644 \u0645\u0631\u062d\u0644\u0629 \u0627\u0644\u0643\u0648\u062a\u0634\u0646\u062c"
        " \u0648\u0627\u0644\u062a\u0649 \u0645\u0646 \u0627\u0644\u0645\u0642\u0631\u0631 \u0628\u062f\u0623\u0647\u0627 \u064a\u0648\u0645 {start_day} \u0627\u0644\u0645\u0648\u0627\u0641\u0642 {start_date} \u0627\u0644\u0633\u0627\u0639\u0629 {start_time}"
        " \u062d\u062a\u0649 {end_time} \u0648\u062a\u0633\u062a\u0645\u0631 \u0644\u0645\u062f\u0629 {duration_weeks} \u0627\u0633\u0627\u0628\u064a\u0639 \u0628\u0645\u0639\u062f\u0644 Session \u0648\u0627\u062d\u062f\u0647 \u0627\u0633\u0628\u0648\u0639\u064a\u0627\u064b"
        " \u0641 \u0646\u0641\u0633 \u0627\u0644\u064a\u0648\u0645 \u0648 \u0627\u0644\u0645\u0648\u0639\u062f \u0627\u0644\u062b\u0627\u0628\u062a \u0648\u0627\u0644\u0645\u062d\u062f\u062f \u0645\u0646 \u0643\u0644 \u0627\u0633\u0628\u0648\u0639\n\n"
        "* \u0628\u0631\u062c\u0627\u0621 \u0627\u0644\u0625\u0646\u0636\u0645\u0627\u0645 \u0644\u062c\u0631\u0648\u0628 \u0627\u0644\u0648\u0627\u062a\u0633\u0627\u0628 \u0627\u0644\u062e\u0627\u0635 \u0628\u0627\u0644\u0645\u062d\u0627\u0636\u0631\u0627\u062a \u0645\u0646 \u062e\u0644\u0627\u0644 \u0627\u0644\u0631\u0627\u0628\u0637 \u0627\u0644\u062a\u0627\u0644\u0649:\n"
        "{link}\n\n"
        "\u0648\u0633\u064a\u062a\u0645 \u0627\u0631\u0633\u0627\u0644 \u062c\u0645\u064a\u0639 \u0627\u0644\u062a\u0641\u0627\u0635\u064a\u0644 \u0627\u0644\u062e\u0627\u0635\u0629 \u0628\u0627\u0644\u062c\u0631\u0648\u0628 \u062e\u0644\u0627\u0644 \u0645\u062c\u0645\u0648\u0639\u0629 \u0627\u0644\u0648\u0627\u062a\u0633\u0627\u0628 \u0628\u0639\u062f \u062f\u062e\u0648\u0644 \u062c\u0645\u064a\u0639 \u0627\u0644\u0637\u0644\u0627\u0628.\n\n"
        "\u0628\u0627\u0644\u062a\u0648\u0641\u064a\u0642 \u0644\u064a\u0643\u0645 \u062c\u0645\u064a\u0639\u064b\u0627 \u0641\u064a \u0645\u0631\u062d\u0644\u0629 {stage} \u0645\u0639 {program} \u2764\ufe0f"
    ),
    "2. Evening Welcome (Arabic)": (
        "\u0645\u0633\u0627\u0621 \u0627\u0644\u062e\u064a\u0631 \u064a\u0627 {name} \ud83d\udc4b\n\n"
        "\u0628\u0646\u0631\u062d\u0628 \u0628\u064a\u0643 \u0641\u064a \u0628\u0631\u0646\u0627\u0645\u062c *{program}* \u2014 \u0645\u0631\u062d\u0644\u0629 *{stage}* \ud83c\udf89\n"
        "\u0645\u0639\u0627\u0643 {coordinator_name}\u060c \u0627\u0644\u0645\u0646\u0633\u0642 \u0627\u0644\u062e\u0627\u0635 \u0628\u0627\u0644\u0645\u062c\u0645\u0648\u0639\u0629.\n\n"
        "\u0647\u064a\u0628\u062f\u0623 \u0645\u0639\u0627\u0646\u0627 \u064a\u0648\u0645 {start_day} \u0627\u0644\u0645\u0648\u0627\u0641\u0642 {start_date} \u0627\u0644\u0633\u0627\u0639\u0629 {start_time}"
        " \u0648\u0647\u064a\u0633\u062a\u0645\u0631 {duration_weeks} \u0623\u0633\u0627\u0628\u064a\u0639.\n\n"
        "\u0627\u0646\u0636\u0645 \u0644\u0644\u062c\u0631\u0648\u0628 \u0645\u0646 \u0627\u0644\u0631\u0627\u0628\u0637 \u062f\u0647:\n{link}\n\n"
        "\u0628\u0627\u0644\u062a\u0648\u0641\u064a\u0642 \u2764\ufe0f"
    ),
    "3. Short Invite (Arabic)": (
        "\u0623\u0647\u0644\u0627\u064b {name} \ud83d\udc4b\n\n"
        "\u062a\u0645 \u0642\u0628\u0648\u0644\u0643 \u0641\u064a \u0628\u0631\u0646\u0627\u0645\u062c *{program}* \u2014 \u0645\u0631\u062d\u0644\u0629 *{stage}*.\n"
        "\u0627\u0646\u0636\u0645 \u0644\u0644\u062c\u0631\u0648\u0628 \u0627\u0644\u0631\u0633\u0645\u064a:\n{link}\n\n"
        "\u0647\u064a\u0628\u062f\u0623 {start_day} {start_date} \u0627\u0644\u0633\u0627\u0639\u0629 {start_time}.\n"
        "\u0628\u0627\u0644\u062a\u0648\u0641\u064a\u0642! \u2764\ufe0f"
    ),
}

def load_group_invite_templates():
    if os.path.exists(GROUP_INVITE_TEMPLATES_FILE):
        try:
            with open(GROUP_INVITE_TEMPLATES_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict) and data:
                return data
        except Exception:
            pass
    return dict(DEFAULT_GROUP_INVITE_TEMPLATES)

def save_group_invite_templates(templates):
    cleaned = {}
    for k, v in templates.items():
        try:
            cleaned[k] = v.encode('utf-16', 'surrogatepass').decode('utf-16')
        except Exception:
            cleaned[k] = v
    with open(GROUP_INVITE_TEMPLATES_FILE, "w", encoding="utf-8") as f:
        json.dump(cleaned, f, ensure_ascii=False, indent=2)

def load_progress():
    if os.path.exists(GROUP_INVITE_PROGRESS_FILE):
        try:
            with open(GROUP_INVITE_PROGRESS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def save_progress(progress):
    with open(GROUP_INVITE_PROGRESS_FILE, "w", encoding="utf-8") as f:
        json.dump(progress, f, ensure_ascii=False, indent=2)

def clear_progress():
    if os.path.exists(GROUP_INVITE_PROGRESS_FILE):
        os.remove(GROUP_INVITE_PROGRESS_FILE)

def normalize_phone(phone_str):
    if not phone_str: return ""
    p = str(phone_str).strip()
    if p.endswith(".00"): p = p[:-3]
    elif p.endswith(".0"): p = p[:-2]
    p = "".join(c for c in p if c.isdigit() or c == "+")
    if not p: return ""
    if p.startswith("+"): return p
    if len(p) == 11 and p.startswith("0"): return "+20" + p[1:]
    if len(p) == 10 and p.startswith(("10","11","12","15")): return "+20" + p
    if p.startswith("20") and len(p) >= 11: return "+" + p
    if p.startswith("00"): return "+" + p[2:]
    if len(p) in (10, 11) and not p.startswith("+"):
        clean = p[1:] if p.startswith("0") else p
        return "+20" + clean
    return "+" + p if not p.startswith("+") else p

def get_first_name(full_name):
    from template_utils import get_first_name as _gfn
    return _gfn(full_name)

def parse_sheet(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        header_row = row
        break
    if not header_row:
        return [], ["Sheet appears to be empty."]

    def _clean_str(val):
        if val is None:
            return ""
        return str(val).replace("\xa0", " ").strip()

    def _norm_header(val):
        text = _clean_str(val).lower()
        text = re.sub(r"[\u064b-\u065f\u0670ـ]", "", text)
        text = text.translate(str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه"}))
        return re.sub(r"[^a-z0-9\u0600-\u06ff]+", "", text)

    raw_headers = [_clean_str(h) for h in header_row]
    lower_headers = [h.casefold() for h in raw_headers]
    norm_headers = [_norm_header(h) for h in raw_headers]

    def _find_col(candidates, norm_candidates=None):
        for c in candidates:
            c_low = c.casefold()
            if c_low in lower_headers:
                return lower_headers.index(c_low)
        if norm_candidates:
            for nc in norm_candidates:
                for idx, nh in enumerate(norm_headers):
                    if nh == nc or (len(nc) >= 4 and nc in nh):
                        return idx
        return None

    arabic_candidates = [
        "arabic name", "arabic_name", "arabic first name", "arabic firstname",
        "arabic", "arabic name only", "student arabic name", "arabic student name",
        "full name (arabic)", "arabic full name", "full name arabic",
        "الاسم بالعربي", "الاسم بالعربية", "الاسم باللغة العربية", "الاسم باللغه العربيه",
        "اسم الطالب بالعربي", "اسم الطالب باللغة العربية", "اسم الطالب باللغه العربيه",
        "اسم المتدرب بالعربي", "اسم المتدرب باللغة العربية", "اسم المشارك بالعربي",
        "الاسم ثلاثي بالعربي", "الاسم الثلاثي بالعربي", "الاسم رباعي بالعربي",
        "الاسم الرباعي بالعربي", "الاسم بالكامل بالعربي", "الاسم عربي", "الاسم (عربي)",
    ]
    arabic_norm_candidates = [
        "arabicname", "arabicfirstname", "studentarabicname", "arabicstudentname",
        "fullnamearabic", "arabicfullname", "arabic", "arabicnameonly",
        "الاسمبالعربي", "الاسمبالعربيه", "الاسمباللغهالعربيه",
        "اسمالطالببالعربي", "اسمالطالبباللغهالعربيه", "اسمالمتدرببالعربي",
        "الاسمثلاثيبالعربي", "الاسمرباعيبالعربي", "الاسمبالكاملبالعربي", "الاسمعربي"
    ]

    name_candidates = [
        "first name", "firstname", "name", "full name", "fullname",
        "student name", "english name", "first_name", "full_name",
        "الاسم", "الاسم الاول", "اسم الطالب", "اسم المتدرب", "اسم المشارك"
    ]
    name_norm_candidates = [
        "firstname", "name", "fullname", "studentname", "englishname",
        "الاسم", "الاسمالاول", "اسمالطالب", "اسمالمتدرب"
    ]

    phone_candidates = [
        "phone", "mobile", "whatsapp", "phone number", "mobile number",
        "phone no", "mobile no", "cell", "telephone",
        "رقم الهاتف", "رقم الموبايل", "رقم الواتساب", "رقم التليفون",
        "موبايل", "هاتف", "تليفون", "واتساب", "tel"
    ]
    phone_norm_candidates = [
        "phone", "mobile", "whatsapp", "phonenumber", "mobilenumber",
        "رقمالهاتف", "رقمالموبايل", "رقمالواتساب", "رقمالتليفون",
        "موبايل", "هاتف", "تليفون", "واتساب"
    ]

    arabic_name_col = _find_col(arabic_candidates, arabic_norm_candidates)
    name_col = _find_col(name_candidates, name_norm_candidates)
    phone_col = _find_col(phone_candidates, phone_norm_candidates)

    if arabic_name_col is None:
        for idx, nh in enumerate(norm_headers):
            if "arabic" in nh or ("عرب" in nh and any(w in nh for w in ["اسم", "name", "طالب", "متدرب"])):
                arabic_name_col = idx
                break

    errors = []
    if arabic_name_col is None and name_col is None:
        errors.append("Could not find a 'Name' or 'Arabic Name' column.")
    if phone_col is None:
        errors.append("Could not find a 'Phone' column.")
    if errors:
        return [], errors

    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw_name = _clean_str(row[name_col]) if (name_col is not None and name_col < len(row)) else ""
        raw_arabic = _clean_str(row[arabic_name_col]) if (arabic_name_col is not None and arabic_name_col < len(row)) else ""
        raw_phone = _clean_str(row[phone_col]) if (phone_col is not None and phone_col < len(row)) else ""

        if not raw_name and not raw_arabic and not raw_phone:
            continue

        # Priority: Arabic Name column first, fallback to Name column
        chosen_name = raw_arabic if raw_arabic else raw_name
        short_name = get_first_name(chosen_name)

        rows.append({
            "name": short_name,
            "arabic_name": raw_arabic,
            "original_name": raw_name or raw_arabic,
            "phone": normalize_phone(raw_phone),
            "raw_phone": raw_phone,
            "row": row_idx
        })
    return rows, []

def build_message(student, template, global_ctx):
    ctx = dict(global_ctx)
    name = student.get("name", "")
    phone = student.get("phone", "")
    ctx["name"] = name
    ctx["Arabic Name"] = name
    ctx["arabic_name"] = name
    ctx["arabic name"] = name
    ctx["first_name"] = name
    ctx["الاسم"] = name
    ctx["الاسم بالعربي"] = name
    ctx["phone"] = phone
    ctx["Mobile"] = phone
    ctx["mobile"] = phone
    return render_template(template, ctx).strip()

def build_all_messages(rows, template, global_ctx):
    return [{
        "name":      s["name"],
        "phone":     s["phone"],
        "raw_phone": s.get("raw_phone",""),
        "row":       s.get("row", 0),
        "message":   build_message(s, template, global_ctx),
    } for s in rows]

def send_batch(messages, engine=None, min_delay=45, max_delay=90,
               daily_limit=25, skip_phones=None,
               on_progress=None, on_delay=None, on_complete=None, stop_event=None):
    """
    Send messages on a background thread with anti-ban measures.
    on_progress(idx, name, phone, ok, detail)  ok=None means skipped/limit
    on_delay(remaining, total)
    on_complete()
    """
    engine      = engine or MessageEngine(min_delay=min_delay, max_delay=max_delay)
    skip_phones = skip_phones or set()
    stop_event  = stop_event  or threading.Event()
    progress    = load_progress()

    def _worker():
        sent_count = 0
        for idx, msg in enumerate(messages):
            if stop_event.is_set(): break
            if sent_count >= daily_limit:
                if on_progress: on_progress(idx, msg["name"], msg["phone"], None, "daily_limit_reached")
                break
            phone = msg["phone"]
            if phone in skip_phones or progress.get(phone, {}).get("status") == "sent":
                if on_progress: on_progress(idx, msg["name"], phone, None, "skipped_already_sent")
                continue
            if not phone:
                key = str(msg.get("row", idx))
                progress[key] = {"status":"failed","timestamp":_now(),"name":msg["name"],"error":"Invalid phone"}
                save_progress(progress)
                if on_progress: on_progress(idx, msg["name"], phone, False, "Invalid phone")
                continue
            ok, detail = engine.send_message(phone, msg["message"])
            progress[phone] = {"status":"sent" if ok else "failed","timestamp":_now(),"name":msg["name"],"error":"" if ok else detail}
            save_progress(progress)
            if on_progress: on_progress(idx, msg["name"], phone, ok, detail)
            if ok: sent_count += 1
            if idx < len(messages)-1 and not stop_event.is_set():
                delay = random.randint(min_delay, max_delay)
                for sec in range(delay):
                    if stop_event.is_set(): break
                    if on_delay: on_delay(delay-sec, delay)
                    time.sleep(1)
                if on_delay and not stop_event.is_set(): on_delay(0, delay)
        if on_complete: on_complete()

    t = threading.Thread(target=_worker, daemon=True)
    t.start()
    return t

def save_log(log_rows, output_dir=None):
    output_dir = output_dir or os.getcwd()
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, GROUP_INVITE_LOG_FILE)
    headers = ["Timestamp","Name","Phone","Status","Error"]
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Log"
    ws.append(headers)
    for row in log_rows: ws.append([row.get(h,"") for h in headers])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col)+2, 60)
    wb.save(path)
    return path

def _now():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")