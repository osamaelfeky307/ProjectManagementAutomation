import os
from datetime import datetime
from fpdf import FPDF
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

class PDFReport(FPDF):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        if os.path.exists(r"C:\Windows\Fonts\arial.ttf"):
            self.add_font("Arial", "", r"C:\Windows\Fonts\arial.ttf")
            self.add_font("Arial", "B", r"C:\Windows\Fonts\arialbd.ttf")
            self.add_font("Arial", "I", r"C:\Windows\Fonts\ariali.ttf")
            self.df = "Arial"
        else:
            self.df = "helvetica"

    def header(self):
        self.set_font(self.df, 'B', 15)
        self.cell(0, 10, 'PM Automation - Analysis Report', border=False, align='C')
        self.ln(20)

    def footer(self):
        self.set_y(-15)
        self.set_font(self.df, 'I', 8)
        self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')

def generate_pdf_report(analysis, totals, conflicts, is_detailed, filepath):
    pdf = PDFReport()
    pdf.add_page()
    df = pdf.df
    
    # Header Information
    pdf.set_font(df, "B", 12)
    pdf.cell(0, 10, "Report Date: ", ln=False)
    pdf.set_font(df, "", 12)
    pdf.cell(0, 10, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), ln=True)
    pdf.ln(5)

    # Analysis Section
    pdf.set_font(df, "B", 14)
    pdf.cell(0, 10, "1. Analysis", ln=True)
    pdf.set_font(df, "", 12)
    
    total_coordinators = len(analysis)
    total_slots = totals.get("Total Slots", 0)
    total_amount = totals.get("Total Amount (EGP)", 0)
    
    pdf.cell(0, 8, f"Total Coordinators: {total_coordinators}", ln=True)
    pdf.cell(0, 8, f"Total Slots: {total_slots}", ln=True)
    pdf.cell(0, 8, f"Total Amount: {total_amount:,.0f} EGP", ln=True)
    pdf.ln(5)

    # Conclusion Section
    pdf.set_font(df, "B", 14)
    pdf.cell(0, 10, "2. Conclusion", ln=True)
    pdf.set_font(df, "", 12)
    
    coord_conflicts = conflicts.get("coordinator_conflicts", [])
    room_conflicts = conflicts.get("room_conflicts", [])
    total_conflicts = len(coord_conflicts) + len(room_conflicts)
    
    if total_conflicts == 0:
        pdf.set_text_color(0, 128, 0) # Green
        pdf.cell(0, 8, "No conflicts detected. All assignments are clean.", ln=True)
        pdf.set_text_color(0, 0, 0)
    else:
        pdf.set_text_color(255, 0, 0) # Red
        pdf.cell(0, 8, f"Warning: {total_conflicts} conflicts detected requiring attention.", ln=True)
        pdf.set_text_color(0, 0, 0)
    pdf.ln(10)

    # More Information Section
    if is_detailed:
        pdf.set_font(df, "B", 14)
        pdf.cell(0, 10, "3. More Information", ln=True)
        
        # Conflict Details
        if total_conflicts > 0:
            pdf.set_font(df, "B", 12)
            pdf.cell(0, 8, "Conflict Details:", ln=True)
            pdf.set_font(df, "", 10)
            if coord_conflicts:
                pdf.cell(0, 8, "Coordinator Conflicts:", ln=True)
                for c in coord_conflicts:
                    centers = ", ".join(c.get("Centers", []))
                    txt = f"- {c.get('Coordinator', '')} is assigned to centers [{centers}] on {c.get('Day', '')} at {c.get('Time', '')}"
                    pdf.multi_cell(0, 6, txt)
            if room_conflicts:
                pdf.cell(0, 8, "Room Conflicts:", ln=True)
                for c in room_conflicts:
                    groups = ", ".join(c.get("Groups", []))
                    txt = f"- Room {c.get('Room Code', '')} in Center {c.get('Center Code', '')} has groups [{groups}] on {c.get('Day', '')} at {c.get('Time', '')}"
                    pdf.multi_cell(0, 6, txt)
            pdf.ln(5)

        # Full Table
        pdf.set_font(df, "B", 12)
        pdf.cell(0, 10, "Coordinator Slots & Financial Dues", ln=True)
        
        # Table Header
        pdf.set_font(df, "B", 10)
        col_widths = [10, 80, 25, 25, 25, 25]
        headers = ["#", "Coordinator", "Main", "Backup", "Total", "Amount"]
        for width, header in zip(col_widths, headers):
            pdf.cell(width, 8, header, border=1, align="C")
        pdf.ln()

        # Table Rows
        pdf.set_font(df, "", 9)
        for i, entry in enumerate(analysis):
            pdf.cell(col_widths[0], 8, str(i + 1), border=1, align="C")
            pdf.cell(col_widths[1], 8, str(entry.get("Coordinator", "")), border=1)
            pdf.cell(col_widths[2], 8, str(entry.get("Main Slots", 0)), border=1, align="C")
            pdf.cell(col_widths[3], 8, str(entry.get("Backup Slots", 0)), border=1, align="C")
            pdf.cell(col_widths[4], 8, str(entry.get("Total Slots", 0)), border=1, align="C")
            pdf.cell(col_widths[5], 8, f"{entry.get('Amount (EGP)', 0):,.0f}", border=1, align="R")
            pdf.ln()
            
    pdf.output(filepath)


def generate_xlsx_report(analysis, totals, conflicts, is_detailed, filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Analysis Report"
    
    # Styles
    bold_font = Font(bold=True)
    header_font = Font(bold=True, size=14)
    red_font = Font(color="FF0000", bold=True)
    green_font = Font(color="008000", bold=True)
    fill_gray = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    row_num = 1
    
    # Header
    ws.cell(row=row_num, column=1, value="PM Automation - Analysis Report").font = header_font
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=6)
    row_num += 2
    
    ws.cell(row=row_num, column=1, value="Report Date:").font = bold_font
    ws.cell(row=row_num, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    row_num += 2
    
    # Analysis
    ws.cell(row=row_num, column=1, value="1. Analysis").font = header_font
    row_num += 1
    
    total_coordinators = len(analysis)
    total_slots = totals.get("Total Slots", 0)
    total_amount = totals.get("Total Amount (EGP)", 0)
    
    ws.cell(row=row_num, column=1, value="Total Coordinators:").font = bold_font
    ws.cell(row=row_num, column=2, value=total_coordinators)
    row_num += 1
    ws.cell(row=row_num, column=1, value="Total Slots:").font = bold_font
    ws.cell(row=row_num, column=2, value=total_slots)
    row_num += 1
    ws.cell(row=row_num, column=1, value="Total Amount (EGP):").font = bold_font
    ws.cell(row=row_num, column=2, value=total_amount)
    row_num += 2

    # Conclusion
    ws.cell(row=row_num, column=1, value="2. Conclusion").font = header_font
    row_num += 1
    
    coord_conflicts = conflicts.get("coordinator_conflicts", [])
    room_conflicts = conflicts.get("room_conflicts", [])
    total_conflicts = len(coord_conflicts) + len(room_conflicts)
    
    if total_conflicts == 0:
        cell = ws.cell(row=row_num, column=1, value="No conflicts detected. All assignments are clean.")
        cell.font = green_font
    else:
        cell = ws.cell(row=row_num, column=1, value=f"Warning: {total_conflicts} conflicts detected requiring attention.")
        cell.font = red_font
    row_num += 2
    
    if is_detailed:
        ws.cell(row=row_num, column=1, value="3. More Information").font = header_font
        row_num += 1
        
        if total_conflicts > 0:
            ws.cell(row=row_num, column=1, value="Conflict Details:").font = bold_font
            row_num += 1
            if coord_conflicts:
                ws.cell(row=row_num, column=1, value="Coordinator Conflicts:").font = bold_font
                row_num += 1
                for c in coord_conflicts:
                    centers = ", ".join(c.get("Centers", []))
                    txt = f"- {c.get('Coordinator', '')} is assigned to centers [{centers}] on {c.get('Day', '')} at {c.get('Time', '')}"
                    ws.cell(row=row_num, column=1, value=txt)
                    row_num += 1
            if room_conflicts:
                ws.cell(row=row_num, column=1, value="Room Conflicts:").font = bold_font
                row_num += 1
                for c in room_conflicts:
                    groups = ", ".join(c.get("Groups", []))
                    txt = f"- Room {c.get('Room Code', '')} in Center {c.get('Center Code', '')} has groups [{groups}] on {c.get('Day', '')} at {c.get('Time', '')}"
                    ws.cell(row=row_num, column=1, value=txt)
                    row_num += 1
            row_num += 1

        ws.cell(row=row_num, column=1, value="Coordinator Slots & Financial Dues").font = header_font
        row_num += 1
        
        headers = ["#", "Coordinator", "Main Slots", "Backup Slots", "Total Slots", "Amount (EGP)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col, value=header)
            cell.font = bold_font
            cell.fill = fill_gray
        row_num += 1
        
        for i, entry in enumerate(analysis):
            ws.cell(row=row_num, column=1, value=i+1)
            ws.cell(row=row_num, column=2, value=entry.get("Coordinator", ""))
            ws.cell(row=row_num, column=3, value=entry.get("Main Slots", 0))
            ws.cell(row=row_num, column=4, value=entry.get("Backup Slots", 0))
            ws.cell(row=row_num, column=5, value=entry.get("Total Slots", 0))
            ws.cell(row=row_num, column=6, value=entry.get("Amount (EGP)", 0))
            row_num += 1

    # Auto-adjust column widths
    for col_idx in range(1, ws.max_column + 1):
        column = get_column_letter(col_idx)
        max_length = 0
        for cell in ws[column]:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        if adjusted_width > 50:
            adjusted_width = 50
        ws.column_dimensions[column].width = adjusted_width

    wb.save(filepath)
