import tempfile
import openpyxl
import pytest
from template_utils import get_first_name
from group_invite_logic import parse_sheet, build_message, build_all_messages

def test_get_first_name_arabic_single():
    assert get_first_name("أحمد") == "أحمد"
    assert get_first_name("سارة") == "سارة"
    assert get_first_name("محمد") == "محمد"

def test_get_first_name_arabic_multiple():
    assert get_first_name("أحمد علي") == "أحمد"
    assert get_first_name("احمد علي السيد") == "احمد"
    assert get_first_name("محمود إبراهيم محمد") == "محمود"

def test_get_first_name_arabic_compounds():
    # Abd compounds
    assert get_first_name("عبد الرحمن أحمد") == "عبد الرحمن"
    assert get_first_name("عبد الله محمد") == "عبد الله"
    assert get_first_name("عبد العزيز علي") == "عبد العزيز"
    assert get_first_name("عبد الفتاح السيسي") == "عبد الفتاح"
    assert get_first_name("عبدالرحمن علي") == "عبدالرحمن"

    # Al-Deen compounds
    assert get_first_name("سيف الدين محمود") == "سيف الدين"
    assert get_first_name("نور الدين علي") == "نور الدين"
    assert get_first_name("صلاح الدين حسن") == "صلاح الدين"

    # Allah compounds
    assert get_first_name("منة الله خالد") == "منة الله"
    assert get_first_name("منه الله خالد") == "منه الله"
    assert get_first_name("هبة الله سعيد") == "هبة الله"
    assert get_first_name("آية الله محمد") == "آية الله"

    # Fatma Al-Zahraa
    assert get_first_name("فاطمة الزهراء علي") == "فاطمة الزهراء"
    assert get_first_name("فاطمه الزهراء علي") == "فاطمه الزهراء"

    # Abu / Om / Ibn
    assert get_first_name("أبو بكر محمد") == "أبو بكر"
    assert get_first_name("ابو بكر محمد") == "ابو بكر"
    assert get_first_name("أم كلثوم علي") == "أم كلثوم"
    assert get_first_name("ابن خلدون") == "ابن خلدون"

def test_get_first_name_english():
    assert get_first_name("Ahmed Ali") == "Ahmed"
    assert get_first_name("Sara Mohamed") == "Sara"

def test_parse_sheet_prefers_arabic_name_over_english_two_names():
    # Create workbook with both English Name and Arabic Name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["First Name", "Arabic Name", "Phone"])
    ws.append(["Ahmed Ali", "أحمد علي محمد", "01012345678"])
    ws.append(["Sara Hassan", "سارة حسن إبراهيم", "01123456789"])
    ws.append(["Abdelrahman Tarek", "عبد الرحمن طارق", "01234567890"])

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name
    wb.save(tmp_path)

    rows, errors = parse_sheet(tmp_path)
    assert not errors, f"Unexpected errors: {errors}"
    assert len(rows) == 3
    # Check that it extracted the first Arabic name only, NOT English two names!
    assert rows[0]["name"] == "أحمد"
    assert rows[0]["phone"] == "+201012345678"
    assert rows[1]["name"] == "سارة"
    assert rows[2]["name"] == "عبد الرحمن"

def test_parse_sheet_arabic_name_column_only():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Arabic Name", "Phone Number"])
    ws.append(["أحمد علي", "01098765432"])

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name
    wb.save(tmp_path)

    rows, errors = parse_sheet(tmp_path)
    assert not errors
    assert len(rows) == 1
    assert rows[0]["name"] == "أحمد"
    assert rows[0]["phone"] == "+201098765432"

def test_parse_sheet_arabic_header_variants():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["الاسم باللغة العربية", "رقم الموبايل"])
    ws.append(["سيف الدين خالد", "01011112222"])

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name
    wb.save(tmp_path)

    rows, errors = parse_sheet(tmp_path)
    assert not errors
    assert len(rows) == 1
    assert rows[0]["name"] == "سيف الدين"

def test_build_message_with_arabic_name():
    student = {"name": "أحمد", "phone": "+201012345678"}
    template = "صباح الخير يا {name}، بنرحب بيكم طلاب {program} فى مرحلة الـ *{stage}*"
    global_ctx = {"program": "NTI", "stage": "Coaching"}
    msg = build_message(student, template, global_ctx)
    assert "صباح الخير يا أحمد،" in msg
    assert "Ahmed Ali" not in msg

def test_build_message_with_arabic_name_placeholder_alias():
    student = {"name": "منى", "phone": "+201151073795"}
    template = "مساء الخير يا {Arabic Name}، بنرحب بيكم طلاب {program} فى مرحلة الـ *{stage}*"
    global_ctx = {"program": "NTI - Creative (i-lance)", "stage": "Coaching"}
    msg = build_message(student, template, global_ctx)
    assert "مساء الخير يا منى،" in msg
