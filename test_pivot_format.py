"""Test pivot-table format detection and unpivoting in coordinator_parser."""
import os
import sys
import openpyxl

# Ensure project root is on path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from coordinator_parser import (
    load_data, validate_columns, _detect_pivot_format, _normalize_header,
)

TEST_DIR = os.path.dirname(os.path.abspath(__file__))
PIVOT_FILE = os.path.join(TEST_DIR, "_test_pivot.xlsx")
STANDARD_FILE = os.path.join(TEST_DIR, "_test_standard.xlsx")
COORD_DIR_FILE = os.path.join(TEST_DIR, "_test_coord_directory.xlsx")


def _create_pivot_file():
    """Create a small pivot-format Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sessions"
    headers = [
        "Row Labels", "Level", "City", "Area", "Center Code", "Center Name", "Time",
        "Sat Code", "Sat Name", "Sun Code", "Sun Name", "Mon Code", "Mon Name",
        "Tue Code", "Tue Name", "Wed Code", "Wed Name",
    ]
    ws.append(headers)
    # Row 1: has Sat, Sun, Mon coordinators
    ws.append(["GRP-001", "Level 1", "Cairo", "Nasr City", "C01", "Center Alpha", "10:00 AM",
               "CO-100", "Ahmed Ali", "CO-200", "Sara Hassan", "CO-300", "Mohamed Youssef",
               "", "", "", ""])
    # Row 2: has Sun and Wed coordinators only
    ws.append(["GRP-002", "Level 2", "Giza", "Haram", "C02", "Center Beta", "02:00 PM",
               "", "", "CO-200", "Sara Hassan", "", "",
               "", "", "CO-400", "Laila Omar"])
    # Row 3: completely empty coordinator columns (should be skipped)
    ws.append(["GRP-003", "Level 1", "Alex", "Sidi Gaber", "C03", "Center Gamma", "04:00 PM",
               "", "", "", "", "", "",
               "", "", "", ""])
    wb.save(PIVOT_FILE)
    print(f"Created pivot test file: {PIVOT_FILE}")


def _create_standard_file():
    """Create a small standard-format Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sessions"
    headers = [
        "Group Code", "Room Code", "Day", "Time", "Center Name",
        "Center Code", "Area", "City", "Main Coordinator", "Phone Number",
    ]
    ws.append(headers)
    ws.append(["GRP-001", "R1", "Saturday", "10:00 AM", "Center Alpha",
               "C01", "Nasr City", "Cairo", "Ahmed Ali", "+201012345678"])
    ws.append(["GRP-002", "R2", "Sunday", "02:00 PM", "Center Beta",
               "C02", "Haram", "Giza", "Sara Hassan", "+201098765432"])
    wb.save(STANDARD_FILE)
    print(f"Created standard test file: {STANDARD_FILE}")


def _create_coordinator_directory():
    """Create a coordinator directory file with names and phones."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Coordinators"
    ws.append(["Coordinator Name", "Phone Number", "Gender"])
    ws.append(["Ahmed Ali", "+201012345678", "Male"])
    ws.append(["Sara Hassan", "+201098765432", "Female"])
    ws.append(["Mohamed Youssef", "+201112223334", "Male"])
    ws.append(["Laila Omar", "+201555666777", "Female"])
    wb.save(COORD_DIR_FILE)
    print(f"Created coordinator directory file: {COORD_DIR_FILE}")


def test_pivot_detection():
    """Test that pivot headers are correctly detected."""
    print("\n=== Test: Pivot Format Detection ===")
    pivot_headers = [
        "Row Labels", "Level", "City", "Area", "Center Code", "Center Name", "Time",
        "Sat Code", "Sat Name", "Sun Code", "Sun Name", "Mon Code", "Mon Name",
    ]
    result = _detect_pivot_format(pivot_headers)
    assert result is not None, "Should detect pivot format"
    assert len(result) >= 2, f"Should find at least 2 day pairs, found {len(result)}"
    day_names = [r[0] for r in result]
    assert "Saturday" in day_names, "Should detect Saturday"
    assert "Sunday" in day_names, "Should detect Sunday"
    assert "Monday" in day_names, "Should detect Monday"
    print(f"  OK Detected {len(result)} day pairs: {day_names}")

    # Standard headers should NOT be detected as pivot
    standard_headers = ["Group Code", "Room Code", "Day", "Time", "Center Name", "Main Coordinator"]
    result2 = _detect_pivot_format(standard_headers)
    assert result2 is None, "Should NOT detect standard format as pivot"
    print("  OK Standard headers correctly rejected as non-pivot")


def test_pivot_parsing():
    """Test end-to-end pivot file parsing and unpivoting."""
    print("\n=== Test: Pivot File Parsing ===")
    _create_pivot_file()
    _create_coordinator_directory()

    data = load_data(PIVOT_FILE, COORD_DIR_FILE)
    rows = validate_columns(data)

    print(f"  Input: 3 pivot rows (1 fully empty)")
    print(f"  Output: {len(rows)} canonical rows")

    # Row 1 has 3 day entries (Sat, Sun, Mon), Row 2 has 2 (Sun, Wed), Row 3 has 0
    assert len(rows) == 5, f"Expected 5 unpivoted rows, got {len(rows)}"
    print("  OK Correct number of unpivoted rows (5)")

    # Check that days are set correctly
    days = [r["day"] for r in rows]
    assert "Saturday" in days, "Should have Saturday row"
    assert "Sunday" in days, "Should have Sunday row"
    assert "Monday" in days, "Should have Monday row"
    assert "Wednesday" in days, "Should have Wednesday row"
    print(f"  OK Days found: {sorted(set(days))}")

    # Check coordinator names
    for row in rows:
        assert row["main_coordinator"], f"Every row should have a coordinator name, got: {row}"
    print("  OK All rows have coordinator names")

    # Check that shared fields are propagated
    grp1_rows = [r for r in rows if r["group_code"] == "GRP-001"]
    assert len(grp1_rows) == 3, f"GRP-001 should have 3 rows, got {len(grp1_rows)}"
    for r in grp1_rows:
        assert r["center_name"] == "Center Alpha", f"Center should be propagated, got: {r['center_name']}"
        assert r["city"] == "Cairo"
        assert r["area"] == "Nasr City"
    print("  OK Shared fields correctly propagated across day rows")

    # Check phone enrichment from directory
    ahmed_rows = [r for r in rows if r["main_coordinator"] == "Ahmed Ali"]
    assert ahmed_rows, "Should have Ahmed Ali rows"
    assert ahmed_rows[0]["main_phone"] == "+201012345678", f"Phone should come from directory, got: {ahmed_rows[0]['main_phone']}"
    print("  OK Phone numbers enriched from coordinator directory")

    # Check gender enrichment
    sara_rows = [r for r in rows if r["main_coordinator"] == "Sara Hassan"]
    assert sara_rows, "Should have Sara Hassan rows"
    assert sara_rows[0]["main_gender"] == "Female", f"Gender should come from directory, got: {sara_rows[0]['main_gender']}"
    print("  OK Gender enriched from coordinator directory")

    print("\n  --- Sample output rows ---")
    for i, row in enumerate(rows):
        print(f"  [{i+1}] {row['group_code']} | {row['day']} | {row['main_coordinator']} (code: {row['main_code']}) | {row['center_name']} | phone: {row['main_phone']}")


def test_standard_parsing():
    """Test that standard format still works (regression check)."""
    print("\n=== Test: Standard Format (Regression) ===")
    _create_standard_file()

    data = load_data(STANDARD_FILE)
    rows = validate_columns(data)

    assert len(rows) == 2, f"Expected 2 rows, got {len(rows)}"
    assert rows[0]["group_code"] == "GRP-001"
    assert rows[0]["day"] == "Saturday"
    assert rows[0]["main_coordinator"] == "Ahmed Ali"
    print(f"  OK Standard format: {len(rows)} rows parsed correctly")


def cleanup():
    for f in (PIVOT_FILE, STANDARD_FILE, COORD_DIR_FILE):
        if os.path.exists(f):
            os.remove(f)
            print(f"  Cleaned up: {f}")


if __name__ == "__main__":
    try:
        test_pivot_detection()
        test_pivot_parsing()
        test_standard_parsing()
        print("\nALL TESTS PASSED")
    except AssertionError as e:
        print(f"\nTEST FAILED: {e}")
        raise
    except Exception as e:
        print(f"\nERROR: {e}")
        raise
    finally:
        cleanup()
